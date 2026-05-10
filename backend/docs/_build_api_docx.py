"""One-off builder for API documentation Word file. Run: python docs/_build_api_docx.py"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent / "API-endpoints.docx"

SECTIONS = [
    (
        "Служебные (без префикса /api/v1)",
        [
            ("GET /", 'Редирект на "/docs".'),
            ("GET /health", 'JSON {"status": "ok"} — признак живости сервиса.'),
        ],
    ),
    (
        "Auth — /api/v1/auth",
        [
            ("POST /register", "UserOut, код 201. Новый пользователь: user_id, name, email, email_verified, xp."),
            ("POST /login", "TokenPair. OAuth2-форма (username=email, password): access_token, refresh_token, token_type."),
            ("POST /login/json", "TokenPair. То же по JSON (email, password)."),
            ("POST /refresh", "TokenPair. Новая пара токенов по refresh_token."),
            ("POST /verify-email", "MessageOut. Результат подтверждения почты по токену в теле."),
            ("GET /verify-email", "MessageOut. То же по query-параметру token."),
            ("POST /change-password", "MessageOut. Смена пароля (нужен Bearer access-токен)."),
            ("POST /forgot-password", "MessageOut. Запрос сброса пароля."),
            ("POST /reset-password", "MessageOut. Новый пароль по токену из письма."),
            ("POST /reset-password/validate", "MessageOut. Проверка токена сброса."),
        ],
    ),
    (
        "Users — /api/v1/users",
        [
            ("GET /leaderboard", "LeaderboardTopOut. Топ по XP: entries с rank, name, xp (до 30)."),
            ("GET /me", "UserOut. Текущий пользователь (Bearer)."),
            ("GET /me/xp-rank", "MyXpRankOut. Ранг, имя, xp текущего пользователя."),
            ("GET /{user_id}", "UserOut. Профиль пользователя по id."),
        ],
    ),
    (
        "Course — /api/v1/course",
        [
            ("GET /course или GET /course/", "Массив TopicOut. Весь курс: топики с элементами (уроки, задания)."),
        ],
    ),
    (
        "Progress — /api/v1/progress",
        [
            ("GET /me", "Объект: ключ topicId, значение — список пройденных itemId."),
            ("GET /me/daily-correct-task-streak", "DailyCorrectTaskStreakOut. consecutive_days."),
            ("GET /me/tasks-progress", "MyTasksProgressOut. Прогресс по топикам и суммы."),
            ("GET /{user_id}", "Та же карта прогресса для указанного пользователя."),
            ("POST /", '{"ok": true}. Отметить элемент пройденным (тело: topicId, itemId; Bearer).'),
        ],
    ),
    (
        "Tasks — /api/v1/tasks",
        [
            ("GET /{item_id}/my-reward-xp", "TaskMyRewardXpOut. Персональная награда XP (Bearer)."),
            ("GET /{item_id}", "TaskGetOut. Условие задания."),
            ("GET /{item_id}/single-choice … /match-pairs", "TaskGetOut для конкретного типа."),
            ("POST /{item_id}/submit/...", "SubmitOut: isCorrect, message, rewardXp (Bearer). Типы: single-choice, fill-in-blank, find-the-bug, code-order, match-pairs."),
            ("GET /{item_id}/attempts", "Список TaskAttemptOut. История попыток (Bearer)."),
        ],
    ),
    (
        "Course admin — /api/v1/course/admin",
        [
            ("Все методы", "HTTP Basic (COURSE_ADMIN_EMAIL / COURSE_ADMIN_PASSWORD)."),
            ("POST /replace-entire-course", '{"ok": true, "topicsLoaded": n}. Полная замена курса.'),
            ("POST /replace-with-sharpik-course", "Загрузка курса из sharpik_course.py."),
            ("POST /topics", "Созданный топик: id, title, order."),
            ("DELETE /topics/{topic_id}", '{"ok": true}.'),
            ("POST /items", "Элемент: id, topicId, type, title, order."),
            ("DELETE /items/{item_id}", '{"ok": true}.'),
            ("POST /theory-blocks", '{"id": …}. Блок теории.'),
            ("DELETE /theory-blocks/{block_id}", '{"ok": true}.'),
            ("POST /tasks/single-choice и др.", "{id, itemId, taskType}. Создание задания."),
            ("DELETE /tasks/{item_id}", '{"ok": true}.'),
        ],
    ),
]


def main() -> None:
    doc = Document()
    st = doc.sections[0]
    st.left_margin = Cm(2)
    st.right_margin = Cm(2)
    st.top_margin = Cm(2)
    st.bottom_margin = Cm(2)

    title = doc.add_heading("Learning Platform API", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Описание эндпоинтов и возвращаемых данных")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.size = Pt(12)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Базовый префикс REST API: /api/v1 (кроме служебных GET / и GET /health). "
        "Защищённые методы: заголовок Authorization: Bearer <access_token>. "
        "Админ курса: HTTP Basic. Подробные схемы запросов и ответов — в интерактивной документации /docs."
    )
    note.paragraph_format.space_after = Pt(12)

    for heading, rows in SECTIONS:
        doc.add_heading(heading, level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        hdr = table.rows[0].cells
        hdr[0].text = "Метод и путь"
        hdr[1].text = "Ответ и краткое пояснение"
        for cell in hdr:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
        for path, desc in rows:
            row = table.add_row().cells
            row[0].text = path
            row[1].text = desc
        doc.add_paragraph()

    doc.save(OUT)
    print(f"Written: {OUT}")


if __name__ == "__main__":
    main()
